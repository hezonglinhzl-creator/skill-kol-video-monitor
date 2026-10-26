from __future__ import annotations

import html
import json
import os
import random
import re
import time
import urllib.parse
import urllib.request
import webbrowser
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "outputs"
OUT_XLSX = Path(os.environ.get("VARLA_VIRAL_OUTPUT_XLSX", OUT_DIR / "一键电摩爆款视频搜索结果.xlsx"))
OUT_HTML = Path(os.environ.get("VARLA_VIRAL_OUTPUT_HTML", OUT_DIR / "一键电摩爆款视频搜索结果.html"))
TOTAL_ROWS = 50
PLATFORM_QUOTA = {
    "TikTok": 30,
    "Instagram": 15,
    "YouTube": 5,
}
AUTO_PUBLIC_INSTAGRAM = False

HEADERS = [
    "序号",
    "平台",
    "视频链接",
    "标题/描述",
    "账号/来源",
    "点赞",
    "播放",
    "关键词",
    "爆火原因",
    "电摩改编方向",
    "优先级",
]

EMOTO_KEYWORDS = [
    "electric dirt bike",
    "emoto",
    "electric motorcycle",
    "electric moto",
    "surron",
    "talaria",
    "off road ebike",
    "electric ride",
    "ride pov electric bike",
    "bike check electric bike",
    "electric off road bike",
    "electric dirtbike",
]

PLATFORM_SEARCHES = [
    ("Instagram", "site:instagram.com/reel"),
]


def open_in_chrome(target: str) -> None:
    chrome = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
    if chrome.exists():
        import subprocess

        subprocess.Popen([str(chrome), target], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    else:
        webbrowser.open(target)


def fetch(url: str, timeout: int = 20) -> str:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Accept-Language": "en-US,en;q=0.9",
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as res:
        return res.read().decode("utf-8", errors="replace")


def parse_count(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "").lower()
    word_match = re.search(r"([\d.]+)\s*(million|thousand|billion)", text)
    if word_match:
        number = float(word_match.group(1))
        scale = {"thousand": 1_000, "million": 1_000_000, "billion": 1_000_000_000}[word_match.group(2)]
        return int(number * scale)
    match = re.match(r"^([\d.]+)\s*([KMB])?$", text, re.I)
    if not match:
        digits = re.sub(r"[^\d]", "", text)
        return int(digits) if digits else 0
    number = float(match.group(1))
    unit = (match.group(2) or "").upper()
    scale = {"K": 1_000, "M": 1_000_000, "B": 1_000_000_000}.get(unit, 1)
    return int(number * scale)


def fmt_count(value: int) -> str:
    if value >= 1_000_000:
        return f"{value / 1_000_000:.1f}M".replace(".0M", "M")
    if value >= 1_000:
        return f"{value / 1_000:.1f}K".replace(".0K", "K")
    return str(value) if value else ""


def normalize_tiktok_video(row: dict, keyword: str) -> dict | None:
    author = row.get("author") or {}
    author_id = author.get("unique_id") or author.get("uniqueId") or row.get("author_unique_id")
    video_id = row.get("video_id") or row.get("id") or row.get("aweme_id")
    if not author_id or not video_id:
        return None

    desc = row.get("title") or row.get("desc") or ""
    views = parse_count(row.get("play_count") or row.get("playCount") or row.get("play"))
    likes = parse_count(row.get("digg_count") or row.get("diggCount") or row.get("likes"))
    link = f"https://www.tiktok.com/@{author_id}/video/{video_id}"
    reason, idea = analyze(f"{desc} @{author_id} {link}", keyword, views, likes)
    return {
        "platform": "TikTok",
        "link": link,
        "title": desc,
        "source": f"@{author_id}",
        "likes": likes,
        "views": views,
        "keyword": keyword,
        "reason": reason,
        "idea": idea,
        "priority": "A",
    }


def search_tiktok(keyword: str, count: int = 18, cursor: int = 0) -> list[dict]:
    url = "https://www.tikwm.com/api/feed/search?" + urllib.parse.urlencode(
        {"keywords": keyword, "count": str(count), "cursor": str(cursor)}
    )
    try:
        payload = json.loads(fetch(url))
    except Exception:
        return []

    data = payload.get("data") or {}
    videos = data.get("videos") or data.get("items") or []
    results = []
    for video in videos:
        item = normalize_tiktok_video(video, keyword)
        if not item:
            continue
        if is_priority_a(item["views"], item["likes"], item["title"], keyword):
            results.append(item)
    return results


def bing_rss_search(query: str, count: int = 6) -> list[dict[str, str]]:
    url = "https://www.bing.com/search?" + urllib.parse.urlencode(
        {"q": query, "format": "rss", "count": str(count)}
    )
    try:
        text = fetch(url, timeout=18)
        root = ET.fromstring(text)
    except Exception:
        return []

    items = []
    for item in root.findall(".//item"):
        title = html.unescape(item.findtext("title") or "").strip()
        link = html.unescape(item.findtext("link") or "").strip()
        desc = html.unescape(item.findtext("description") or "").strip()
        if link and "bing.com" not in link:
            items.append({"title": title, "link": link, "desc": desc})
    return items


def bing_html_reel_search(query: str, count: int = 20) -> list[str]:
    url = "https://www.bing.com/search?" + urllib.parse.urlencode({"q": query, "count": str(count)})
    try:
        page = html.unescape(fetch(url, timeout=18))
    except Exception:
        return []

    links = []
    patterns = [
        r"https?://(?:www\.)?instagram\.com/reel/[^?&#\"<> ]+",
        r"https?%3A%2F%2F(?:www\.)?instagram\.com%2Freel%2F[^%?&#\"<> ]+",
    ]
    for pattern in patterns:
        for match in re.findall(pattern, page, flags=re.I):
            link = urllib.parse.unquote(match)
            clean = clean_video_link("Instagram", link)
            if clean and clean not in links:
                links.append(clean)
    return links


def clean_video_link(platform: str, link: str) -> str | None:
    link = html.unescape(link or "").strip()
    if not link:
        return None
    low = link.lower()
    if platform == "Instagram":
        match = re.search(r"https?://(?:www\.)?instagram\.com/reel/[^/?#]+", link, re.I)
        return match.group(0) if match else None
    if platform == "YouTube":
        match = re.search(r"https?://(?:www\.)?youtube\.com/shorts/[^/?#&]+", link, re.I)
        if match:
            return match.group(0)
        if "youtube.com/watch" in low and "v=" in low:
            return link.split("&")[0]
    return None


def infer_source(platform: str, link: str, title: str) -> str:
    if platform == "Instagram":
        return "Instagram Reel"
    if platform == "YouTube":
        return "YouTube Shorts"
    title = re.sub(r"\s*[-|]\s*(TikTok|Instagram|YouTube).*$", "", title, flags=re.I)
    return title[:40]


def search_instagram_youtube(keyword: str) -> list[dict]:
    rows = []
    for platform, operator in PLATFORM_SEARCHES:
        query = f'{operator} "{keyword}"'
        for link in bing_html_reel_search(query):
            reason, idea = analyze(f"{keyword} {link}", keyword, 0, 0)
            rows.append(
                {
                    "platform": "Instagram",
                    "link": link,
                    "title": keyword,
                    "source": "Instagram Reel",
                    "likes": 0,
                    "views": 0,
                    "keyword": keyword,
                    "reason": reason,
                    "idea": idea,
                    "priority": "A",
                }
            )
        for item in bing_rss_search(query, count=6):
            link = clean_video_link(platform, item["link"])
            if not link:
                continue
            title = item["title"] or item["desc"]
            reason, idea = analyze(f"{title} {link}", keyword, 0, 0)
            rows.append(
                {
                    "platform": platform,
                    "link": link,
                    "title": title,
                    "source": infer_source(platform, link, title),
                    "likes": 0,
                    "views": 0,
                    "keyword": keyword,
                    "reason": reason,
                    "idea": idea,
                    "priority": "A",
                }
            )
    return rows


def find_json_objects(node: object, key: str) -> list[dict]:
    found = []
    if isinstance(node, dict):
        if key in node and isinstance(node[key], dict):
            found.append(node[key])
        for value in node.values():
            found.extend(find_json_objects(value, key))
    elif isinstance(node, list):
        for value in node:
            found.extend(find_json_objects(value, key))
    return found


def text_from_runs(value: object) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        if isinstance(value.get("simpleText"), str):
            return value["simpleText"]
        runs = value.get("runs")
        if isinstance(runs, list):
            return "".join(str(run.get("text", "")) for run in runs if isinstance(run, dict))
    return ""


def extract_youtube_initial_data(page: str) -> dict | None:
    match = re.search(r"ytInitialData\s*=\s*({.*?});</script>", page)
    if not match:
        match = re.search(r"var ytInitialData\s*=\s*({.*?});", page)
    if not match:
        return None
    try:
        return json.loads(match.group(1))
    except Exception:
        return None


def search_youtube(keyword: str) -> list[dict]:
    query = f"{keyword} shorts"
    url = "https://www.youtube.com/results?" + urllib.parse.urlencode({"search_query": query})
    try:
        page = fetch(url, timeout=20)
    except Exception:
        return []
    rows = []
    seen_ids = set()

    blocks = re.findall(r'\{"shortsLockupViewModel":(.*?)(?=\{"shortsLockupViewModel"|$)', page, re.S)
    for block in blocks:
        id_match = re.search(r'"videoId":"([A-Za-z0-9_-]{11})"', block)
        text_match = re.search(r'"accessibilityText":"(.*?)"', block, re.S)
        if not id_match or not text_match:
            continue
        video_id = id_match.group(1)
        if not video_id or video_id in seen_ids:
            continue
        try:
            raw_text = json.loads('"' + text_match.group(1).replace('"', '\\"') + '"')
        except Exception:
            raw_text = text_match.group(1)
        raw_text = html.unescape(raw_text)
        seen_ids.add(video_id)
        title = re.sub(r",\s*[\d.,]+\s*(?:thousand|million|billion)?\s*views?\s*-\s*play Short.*$", "", raw_text, flags=re.I)
        views = parse_count(raw_text)
        if not is_priority_a(views, 0, title, keyword):
            continue
        short_link = f"https://www.youtube.com/shorts/{video_id}"
        reason, idea = analyze(f"{title} {short_link}", keyword, views, 0)
        rows.append(
            {
                "platform": "YouTube",
                "link": short_link,
                "title": title,
                "source": "YouTube Shorts",
                "likes": 0,
                "views": views,
                "keyword": keyword,
                "reason": reason,
                "idea": idea,
                "priority": "A",
            }
        )

    return rows


def is_priority_a(views: int, likes: int, title: str, keyword: str) -> bool:
    text = f"{title} {keyword}".lower()
    strong_topic = any(
        word in text
        for word in [
            "surron",
            "talaria",
            "electric dirt",
            "electric motorcycle",
            "emoto",
            "off road",
            "pov",
            "bike check",
        ]
    )
    return views >= 100_000 or likes >= 3_000 or (strong_topic and views >= 30_000)


def normalize_link(link: str) -> str:
    link = (link or "").strip()
    if not link:
        return ""
    parsed = urllib.parse.urlparse(link)
    return urllib.parse.urlunparse((parsed.scheme, parsed.netloc.lower(), parsed.path.rstrip("/"), "", "", ""))


def links_from_workbook(path: Path) -> set[str]:
    if not path.exists():
        return set()
    try:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        links = set()
        for row in range(2, ws.max_row + 1):
            value = normalize_link(str(ws.cell(row, 3).value or ""))
            if value:
                links.add(value)
        return links
    except Exception:
        return set()


def previous_output_links() -> set[str]:
    links = set()
    for path in [OUT_DIR / "一键电摩爆款视频搜索结果.xlsx", OUT_DIR / "viral_video_research.xlsx"]:
        if path.resolve() != OUT_XLSX.resolve():
            links.update(links_from_workbook(path))
    history_dir = OUT_DIR / "history"
    if history_dir.exists():
        for path in history_dir.glob("*.xlsx"):
            if path.resolve() != OUT_XLSX.resolve():
                links.update(links_from_workbook(path))
    return links


def analyze(title: str, keyword: str, views: int, likes: int) -> tuple[str, str]:
    text = f"{title} {keyword}".lower()
    view_label = fmt_count(views) if views else ""
    like_label = fmt_count(likes) if likes else ""
    metric = f"{view_label} 播放" if view_label else (f"{like_label} 点赞" if like_label else "")

    def pack(reason: str, idea: str) -> tuple[str, str]:
        prefix = f"{metric}，" if metric else ""
        return prefix + reason, idea

    if "7512855279214562582" in text or ("unilad" in text and ("unicycle" in text or "future of travel" in text)):
        return pack(
            "这条爆点是擦边/猎奇出行装备组合：人物画面先吸引停留，少见的电动独轮车再制造新鲜感；标题把它包装成 future of travel，评论会围绕好不好看、安不安全、实不实用争论。",
            "Varla 不照搬擦边。改成“路人第一眼看到 Varla 的反应”：第一秒人车同框+回头反应，再切整车外观、灯光/仪表/轮胎近景、街区低速安全通过。",
        )

    if "7115414695728663814" in text or "thekidscars" in text or "little driver" in text or "kidscar" in text or "kids car" in text:
        return pack(
            "这条爆点是小朋友开车/萌娃反差：小孩和驾驶行为形成强反差，观众会因为可爱、惊讶和安全边界停留；核心不是车性能，而是人物设定带来的情绪。",
            "Varla 不做儿童驾驶。改成“家人围观新车/新手第一次安全试骑”：成人低速骑行，孩子或家人只做旁观反应，保留家庭户外氛围和反差感。",
        )

    if any(w in text for w in ["let me know", "what i missed", "comment", "comments", "pros and cons", "top things"]):
        return pack(
            "这条更像口播清单/评论互动型：人和车同框，用字幕重点让观众判断他说得对不对、漏了什么，所以互动来自补充和争论，不是单纯骑行爽感。",
            "Varla 改成“5 things I noticed after riding Varla”：人站车旁口播，每点配一个近景或短骑行镜头，结尾问观众想看续航、越野还是街道体验。",
        )

    if any(w in text for w in ["sale alert", "limited time", "selling fast", "last units", "discount", "price", "$"]):
        return pack(
            "这条是低价/限时促销钩子：用户停留是为了判断真假、值不值、有没有坑；画面通常只是辅助证明有这台车，不是靠骑行动作爆。",
            "Varla 改成真实优惠+卖点展示：先给一句明确利益点，再切整车、轮胎、避震、灯光、仪表近景，最后给官网/优惠码 CTA。",
        )

    if any(w in text for w in ["dad my bike", "stopped working", "reaction", "trolling", "prank", "tag your parents", "mom", "mother", "daughter"]):
        return pack(
            "这条靠家庭反应/轻剧情爆：观众看的是人物关系和反差，不是参数；评论会代入父母、孩子、朋友之间的反应。",
            "Varla 改成“家人第一次看到 Varla”：先拍表情反应，再拍启动、灯光、短距离安全骑行，用生活场景带出产品。",
        )

    if any(w in text for w in ["vs", "which is better", "comparison", "drag race", "stock", "head to head"]):
        return pack(
            "这条靠对比冲突爆：观众会留下来看谁更强、差距在哪里；标题天然有悬念，容易引导评论站队。",
            "Varla 改成安全对比：Varla vs 普通代步车，对比起步、爬小坡、刹停、舒适度，场地用封闭空地或低速街区。",
        )

    if any(w in text for w in ["bike check", "check", "review", "worth", "budget", "best"]):
        return pack(
            "这条是配置判断型：用户想快速知道这车值不值、哪里强、哪里适合自己；爆点来自信息密度和购买判断。",
            "Varla 改成 15 秒 Bike Check：轮胎、电机、避震、灯光、座椅、车架各给一个近景，每个镜头只写一个卖点。",
        )

    if any(w in text for w in ["unbox", "assembly", "setup", "ordered", "arrived", "in!!!"]):
        return pack(
            "这条是到货/组装过程型：观众想知道大件产品收到后麻不麻烦、配件是否完整、第一次启动是否顺利。",
            "Varla 改成开箱四步：外箱到达、配件平铺、关键部件组装、首次启动和短距离试骑，原始素材可二剪成教程。",
        )

    if any(w in text for w in ["pov", "first ride", "ride pov"]):
        return pack(
            "这条靠第一视角代入感：观众像自己在骑，前 3 秒必须直接进入动作；爆点在起步、转弯、路面通过的连续体验。",
            "Varla 改成第一视角脚本：上车按电门、仪表亮起、起步、过弯、停下回头看车；场景用街区慢速、草地小坡或土路通过。",
        )

    if any(w in text for w in ["wheelie", "fastest", "speed", "50mph", "acceleration", "insane", "crazy"]):
        return pack(
            "这条靠速度/刺激/强动作吸引停留，第一秒通常就有动态画面；但品牌改编不能放大危险动作。",
            "Varla 改成可控爽感：低机位拍起步、路面掠过、转弯压弯、远景经过，字幕强调轻越野爽感和安全场地。",
        )

    if any(w in text for w in ["surron", "talaria", "eride", "ultrabee", "dirt bike", "off road", "emoto", "electric motorcycle", "electric dirt"]):
        return pack(
            "这条命中电摩/越野电动两轮受众，相关性高；但如果只看关键词还不够，需要打开视频确认它是骑行、展示、口播还是剧情。",
            "Varla 改成单一方向：骑行爽感、卖点展示、开箱组装三选一，不要把参数、剧情和 POV 混成一条。",
        )

    return pack(
        "这条只能先判断为泛热点，标题和账号不足以确认画面动作；不要直接套 POV 或卖点模板，需要打开视频看具体内容。",
        "Varla 改编先保守：提取开头钩子、人物/产品关系、字幕节奏，再换成安全骑行或静态卖点展示。",
    )


def collect_results() -> list[dict]:
    # A new automatic run must not reuse links from previous result/history files.
    blocked = previous_output_links()
    seen = set(blocked)
    buckets = {"TikTok": [], "Instagram": [], "YouTube": []}

    keywords = EMOTO_KEYWORDS[:]
    random.shuffle(keywords)

    tiktok_cursors = [0, 12, 24, 36, 48, 60, 72, 96, 120, 156, 192]
    random.shuffle(tiktok_cursors)
    for keyword in keywords:
        if len(buckets["TikTok"]) >= 60:
            break
        for cursor in tiktok_cursors:
            for item in search_tiktok(keyword, count=12, cursor=cursor):
                clean = normalize_link(item["link"])
                if clean and clean not in seen:
                    item["link"] = clean
                    seen.add(clean)
                    buckets["TikTok"].append(item)
                    if len(buckets["TikTok"]) >= 60:
                        break
            if len(buckets["TikTok"]) >= 60:
                break
            time.sleep(0.18)

    for keyword in keywords[:8]:
        if len(buckets["YouTube"]) >= 10:
            break
        for item in search_youtube(keyword):
            clean = normalize_link(item["link"])
            if clean and clean not in seen:
                item["link"] = clean
                seen.add(clean)
                buckets["YouTube"].append(item)
                if len(buckets["YouTube"]) >= 10:
                    break
        time.sleep(0.25)

    if AUTO_PUBLIC_INSTAGRAM:
        for keyword in keywords[:3]:
            if len(buckets["Instagram"]) >= PLATFORM_QUOTA["Instagram"]:
                break
            for item in search_instagram_youtube(keyword):
                clean = normalize_link(item["link"])
                if clean and clean not in seen:
                    item["link"] = clean
                    seen.add(clean)
                    buckets["Instagram"].append(item)
                    if len(buckets["Instagram"]) >= PLATFORM_QUOTA["Instagram"]:
                        break
            time.sleep(0.25)

    for platform in buckets:
        buckets[platform].sort(key=lambda x: (x["views"], x["likes"]), reverse=True)

    rows = []
    selected = set()

    def add_items(items: list[dict], limit: int) -> None:
        for item in items:
            if len(rows) >= TOTAL_ROWS or limit <= 0:
                return
            clean = normalize_link(item["link"])
            if not clean or clean in selected:
                continue
            item["link"] = clean
            rows.append(item)
            selected.add(clean)
            limit -= 1

    add_items(buckets["TikTok"], PLATFORM_QUOTA["TikTok"])
    add_items(buckets["Instagram"], PLATFORM_QUOTA["Instagram"])
    add_items(buckets["YouTube"], PLATFORM_QUOTA["YouTube"])

    remaining = []
    for platform in ["TikTok", "Instagram", "YouTube"]:
        remaining.extend([item for item in buckets[platform] if normalize_link(item["link"]) not in selected])
    remaining.sort(key=lambda x: (x["views"], x["likes"]), reverse=True)
    add_items(remaining, TOTAL_ROWS - len(rows))

    rows.sort(key=lambda x: (["TikTok", "Instagram", "YouTube"].index(x["platform"]), -x["views"], -x["likes"]))
    return rows[:TOTAL_ROWS]


def write_xlsx(rows: list[dict]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "A 级爆款视频"
    ws.append(HEADERS)

    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, item in enumerate(rows, start=1):
        ws.append(
            [
                index,
                item["platform"],
                item["link"],
                item["title"],
                item["source"],
                fmt_count(item["likes"]),
                fmt_count(item["views"]),
                item["keyword"],
                item["reason"],
                item["idea"],
                item["priority"],
            ]
        )

    widths = [8, 12, 50, 42, 20, 12, 12, 22, 72, 72, 10]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)


def write_html(rows: list[dict]) -> None:
    def cell(value: object, link: bool = False) -> str:
        value = str(value or "")
        if link and value:
            return f'<td><a href="{html.escape(value)}" target="_blank">打开视频</a></td>'
        return f"<td>{html.escape(value)}</td>"

    trs = []
    for index, item in enumerate(rows, start=1):
        values = [
            index,
            item["platform"],
            item["link"],
            item["title"],
            item["source"],
            fmt_count(item["likes"]),
            fmt_count(item["views"]),
            item["keyword"],
            item["reason"],
            item["idea"],
            item["priority"],
        ]
        tds = [cell(values[0]), cell(values[1]), cell(values[2], True)] + [cell(v) for v in values[3:]]
        trs.append("<tr>" + "".join(tds) + "</tr>")

    OUT_HTML.write_text(
        f"""<!doctype html>
<meta charset="utf-8">
<title>一键电摩爆款视频搜索结果</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;margin:24px;color:#1f2329}}
table{{border-collapse:collapse;width:1800px;font-size:13px;table-layout:fixed}}
th,td{{border:1px solid #d9dce3;padding:8px;vertical-align:top;word-break:break-word;line-height:1.45}}
th{{background:#d9eaf7}}
th:nth-child(9),td:nth-child(9),th:nth-child(10),td:nth-child(10){{width:330px}}
a{{color:#245bdb}}
</style>
<h1>一键电摩爆款视频搜索结果（只保留优先级 A）</h1>
<p>生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}</p>
<p>Excel 文件：{html.escape(str(OUT_XLSX))}</p>
<table>
<tr>{''.join(f'<th>{html.escape(h)}</th>' for h in HEADERS)}</tr>
{''.join(trs)}
</table>
""",
        encoding="utf-8",
    )


def validate_rows(rows: list[dict]) -> list[str]:
    problems = []
    for idx, item in enumerate(rows, 1):
        link = item["link"].lower()
        if item["priority"] != "A":
            problems.append(f"第 {idx} 行不是 A 级")
        if item["platform"] == "TikTok" and not re.search(r"tiktok\.com/@[^/]+/video/\d+", link):
            problems.append(f"第 {idx} 行 TikTok 链接不是视频页：{item['link']}")
        if item["platform"] == "Instagram" and "instagram.com/reel/" not in link:
            problems.append(f"第 {idx} 行 Instagram 链接不是 Reel：{item['link']}")
        if item["platform"] == "YouTube" and "youtube.com/shorts/" not in link:
            problems.append(f"第 {idx} 行 YouTube 不是 Shorts：{item['link']}")
    return problems


def main() -> None:
    print("正在搜索 TikTok / Instagram / YouTube 电摩爆款视频，只保留优先级 A...")
    rows = collect_results()
    problems = validate_rows(rows)
    if not rows:
        print("没有抓到有效视频。请稍后重试，或先登录平台后使用浏览器采集插件。")
        return
    if problems:
        print("自检发现问题：")
        for problem in problems[:10]:
            print(problem)
        print("已停止写入，避免再次混入错误网页。")
        return
    write_xlsx(rows)
    write_html(rows)
    print(f"已找到 {len(rows)} 条 A 级候选视频")
    print(f"Excel: {OUT_XLSX}")
    print(f"HTML: {OUT_HTML}")
    open_in_chrome(OUT_HTML.as_uri())


if __name__ == "__main__":
    main()
