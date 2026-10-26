from __future__ import annotations

import csv
import html
import json
import re
import sys
import urllib.parse
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
XLSX_PATH = OUTPUT_DIR / "viral_video_research.xlsx"
CSV_PATH = OUTPUT_DIR / "viral_video_research.csv"
MD_PATH = OUTPUT_DIR / "viral_video_research_feishu.md"
PORT = 8766

KEYWORDS = [
    "#ebike",
    "#electricbike",
    "#electricbikes",
    "#electricride",
    "#electricdirtbike",
    "#electricoffroadbike",
    "#offroadebike",
    "#emoto",
    "#emotos",
    "#electricmoto",
    "#bikelife",
    "#bikecheck",
    "#ridepov",
]

TREND_SEEDS = [
    "black vs white foam rust removal challenge",
    "lawn mowing race challenge",
    "satisfying pressure wash transformation",
    "before and after cleaning challenge",
    "speed comparison challenge",
    "POV first ride viral trend",
    "this or that product comparison",
    "oddly satisfying repair transformation",
]

HEADERS = [
    "No.",
    "Date",
    "Platform",
    "Video Link",
    "Account",
    "Caption / Title",
    "Likes",
    "Views",
    "Keyword / Trend",
    "Viral Reason",
    "E-Moto Adaptation Idea",
    "Suggested Varla Content",
    "Priority",
    "Status",
    "Notes",
]


def today() -> str:
    return datetime.now().strftime("%Y.%m.%d")


def compact_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()[:500]


def normalize_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        return ""
    parsed = urllib.parse.urlparse(url)
    return urllib.parse.urlunparse((parsed.scheme, parsed.netloc, parsed.path.rstrip("/"), "", "", ""))


def infer_platform(url: str) -> str:
    low = url.lower()
    if "tiktok.com" in low:
        return "TikTok"
    if "instagram.com" in low:
        return "Instagram"
    if "youtube.com" in low or "youtu.be" in low:
        return "YouTube"
    return ""


def infer_account(url: str, title: str, body_text: str = "") -> str:
    text = " ".join([url or "", title or "", body_text or ""])
    patterns = [
        r"tiktok\.com/@([A-Za-z0-9._-]+)",
        r"instagram\.com/([A-Za-z0-9._-]+)/(?:reel|p|tv)?",
        r"@([A-Za-z0-9._-]{2,})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            account = match.group(1).strip()
            if account.lower() not in {"reel", "p", "explore", "shorts", "watch"}:
                return "@" + account.lstrip("@")
    if "youtube.com" in (url or "").lower():
        match = re.search(r" - YouTube$", title or "")
        if match:
            return compact_text((title or "").replace(" - YouTube", ""))
    return ""


def parse_number(text: str) -> float | None:
    if not text:
        return None
    raw = text.replace(",", "").strip().upper()
    match = re.search(r"(\d+(?:\.\d+)?)\s*([KMB万亿]?)", raw)
    if not match:
        return None
    number = float(match.group(1))
    unit = match.group(2)
    if unit == "K":
        number *= 1_000
    elif unit == "M":
        number *= 1_000_000
    elif unit == "B":
        number *= 1_000_000_000
    elif unit == "万":
        number *= 10_000
    elif unit == "亿":
        number *= 100_000_000
    return number


def compact_number(value: str) -> str:
    number = parse_number(value)
    if number is None:
        return compact_text(value)
    if number >= 1_000_000:
        return f"{number / 1_000_000:.1f}".rstrip("0").rstrip(".") + "M"
    if number >= 1_000:
        decimals = 0 if number >= 100_000 else 1
        return f"{number / 1_000:.{decimals}f}".rstrip("0").rstrip(".") + "K"
    return str(int(round(number)))


def detect_keyword(text: str) -> str:
    low = (text or "").lower()
    hits = [kw for kw in KEYWORDS if kw.lower().replace("#", "") in low or kw.lower() in low]
    if hits:
        return ", ".join(hits[:4])
    for trend in TREND_SEEDS:
        tokens = [t for t in re.split(r"\W+", trend.lower()) if len(t) > 3]
        if sum(1 for t in tokens if t in low) >= 2:
            return trend
    return ""


def analyze_video(title: str, body_text: str, keyword: str) -> tuple[str, str, str, str]:
    text = " ".join([title or "", body_text or "", keyword or ""]).lower()
    reasons: list[str] = []
    ideas: list[str] = []
    suggested = "骑行爽感类 / POV Ride"
    priority = "B"

    if any(w in text for w in ["pov", "first ride", "ride", "acceleration", "speed", "fast"]):
        reasons.append("POV/速度感强，用户代入感高，前几秒容易抓注意力。")
        ideas.append("改成 Varla 第一视角起步加速、草地/土路/安全街道骑行，前 3 秒直接给动作镜头。")
        suggested = "骑行爽感类 / POV Ride"
        priority = "A"
    if any(w in text for w in ["check", "feature", "review", "which is better", "comparison", "vs"]):
        reasons.append("有对比或测评结构，用户容易停留看结论。")
        ideas.append("改成 Varla vs 普通 e-bike / gas mini bike 的卖点对比，突出电机、轮胎、减震、续航。")
        suggested = "卖点展示类 / Feature Showcase"
        priority = "A"
    if any(w in text for w in ["unbox", "assembly", "setup", "package"]):
        reasons.append("开箱/组装能降低购买顾虑，适合转化。")
        ideas.append("改成 Varla 收货、拆箱、装车、第一次启动、第一次试骑流程。")
        suggested = "开箱组装类 / Unboxing"
    if any(w in text for w in ["foam", "rust", "clean", "pressure wash", "satisfying", "before and after"]):
        reasons.append("清洁/修复/前后对比有强烈视觉变化，适合 satisfying 爆款逻辑。")
        ideas.append("改成泥地骑完后的 Varla 清洗、泡沫覆盖、前后对比，做电摩版 satisfying transformation。")
        suggested = "热点改编 / Cleaning Transformation"
        priority = "A"
    if any(w in text for w in ["mowing", "lawn", "race", "challenge", "competition"]):
        reasons.append("比赛/挑战结构简单，观众容易理解并评论站队。")
        ideas.append("改成电摩骑行挑战：谁更快完成路线、谁能更稳通过草地/土路，强调安全和合法场地。")
        suggested = "热点改编 / Challenge"
        priority = "A"
    if any(w in text for w in ["funny", "fail", "reaction", "crazy"]):
        reasons.append("反应/反差/轻剧情容易触发评论和转发。")
        ideas.append("改成朋友第一次试骑 Varla 的真实反应，保留轻松搞笑但不危险。")
        suggested = "UGC 反应类 / First Ride Reaction"

    if not reasons:
        reasons.append("可能依靠清晰主题、短时长、强画面或账号粉丝基础获得传播。")
        ideas.append("可先复刻结构：3 秒钩子 + 1 个明确卖点/动作 + 结尾 CTA。")

    return "；".join(reasons[:2]), "；".join(ideas[:2]), suggested, priority


def ensure_workbook() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if XLSX_PATH.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "爆款视频库"
    ws.append(HEADERS)
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [6, 12, 12, 38, 18, 45, 12, 12, 24, 42, 48, 28, 10, 12, 30]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    wb.save(XLSX_PATH)


def load_existing_urls(ws) -> set[str]:
    urls = set()
    for row in range(2, ws.max_row + 1):
        url = normalize_url(str(ws.cell(row, 4).value or ""))
        if url:
            urls.add(url)
    return urls


def get_records_sheet(wb):
    if "爆款视频库" in wb.sheetnames:
        return wb["爆款视频库"]
    return wb.active


def save_record(params: dict[str, str]) -> tuple[int, dict[str, str]]:
    ensure_workbook()
    url = normalize_url(params.get("url", ""))
    title = compact_text(params.get("title", ""))
    body_text = compact_text(params.get("body_text", ""))
    platform = params.get("platform") or infer_platform(url)
    account = params.get("account") or infer_account(url, title, body_text)
    likes = compact_number(params.get("likes", ""))
    views = compact_number(params.get("views", ""))
    keyword = params.get("keyword") or detect_keyword(" ".join([title, body_text, url]))
    viral_reason, idea, suggested, priority = analyze_video(title, body_text, keyword)

    if not url:
        raise ValueError("Missing video URL.")
    if not platform:
        raise ValueError("Only TikTok, Instagram, and YouTube links are supported.")

    wb = load_workbook(XLSX_PATH)
    ws = get_records_sheet(wb)
    existing = load_existing_urls(ws)
    if url in existing:
        for row in range(2, ws.max_row + 1):
            if normalize_url(str(ws.cell(row, 4).value or "")) == url:
                return row, {"status": "duplicate", "url": url}

    row = ws.max_row + 1
    values = [
        row - 1,
        today(),
        platform,
        url,
        account,
        title or body_text[:120],
        likes,
        views,
        keyword,
        viral_reason,
        idea,
        suggested,
        priority,
        "待复盘",
        compact_text(params.get("notes", "")),
    ]
    ws.append(values)
    for cell in ws[row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(XLSX_PATH)
    write_csv_and_md()
    return row, {"status": "saved", "url": url, "account": account, "platform": platform}


def records_from_workbook() -> list[list[str]]:
    ensure_workbook()
    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = get_records_sheet(wb)
    rows: list[list[str]] = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        rows.append([str(cell or "") for cell in row])
    return rows


def write_csv_and_md() -> None:
    rows = records_from_workbook()
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows)
    with MD_PATH.open("w", encoding="utf-8") as f:
        f.write("# 外网爆款视频库\n\n")
        f.write("| " + " | ".join(HEADERS) + " |\n")
        f.write("|" + "|".join(["---"] * len(HEADERS)) + "|\n")
        for row in rows[1:]:
            clean = [str(x).replace("\n", " ").replace("|", "/") for x in row]
            f.write("| " + " | ".join(clean) + " |\n")


def search_urls() -> list[tuple[str, str, str]]:
    links: list[tuple[str, str, str]] = []
    queries = KEYWORDS + TREND_SEEDS
    for query in queries:
        clean = query.replace("#", "").strip()
        encoded = urllib.parse.quote(clean)
        links.append(("TikTok", query, f"https://www.tiktok.com/search/video?q={encoded}"))
        links.append(("Instagram", query, f"https://www.instagram.com/explore/search/keyword/?q={encoded}"))
        links.append(("YouTube", query, f"https://www.youtube.com/results?search_query={encoded}"))
    return links


def page(body: str) -> bytes:
    return f"""<!doctype html>
<meta charset="utf-8">
<title>Viral Video Finder</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;max-width:1100px;margin:28px auto;color:#1f2329;line-height:1.5}}
table{{border-collapse:collapse;width:100%;font-size:13px}}
td,th{{border:1px solid #d9dce3;padding:8px;vertical-align:top}}
th{{background:#e8f0fe}}
.ok{{background:#e8f7ee;border:1px solid #b7ebc6;padding:12px;border-radius:8px}}
.err{{background:#fff1f0;border:1px solid #ffa39e;padding:12px;border-radius:8px}}
a.button{{display:inline-block;background:#245bdb;color:#fff;text-decoration:none;padding:8px 12px;border-radius:6px;margin:4px}}
code{{background:#f5f6f8;padding:2px 4px;border-radius:4px}}
</style>
{body}
""".encode("utf-8")


class Handler(BaseHTTPRequestHandler):
    def send_html(self, body: str) -> None:
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(page(body))

    def send_json(self, data: dict) -> None:
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False).encode("utf-8"))

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        params = {k: v[0] for k, v in urllib.parse.parse_qs(parsed.query).items()}
        if parsed.path == "/health":
            self.send_json({"ok": True})
            return
        if parsed.path == "/add":
            try:
                row, data = save_record(params)
                self.send_html(
                    f'<div class="ok">已记录到第 {row} 行：{html.escape(data.get("platform",""))} '
                    f'{html.escape(data.get("account",""))}</div><p>可以关闭本页，继续采集下一个视频。</p>'
                )
            except Exception as exc:
                self.send_html(f'<div class="err">记录失败：{html.escape(str(exc))}</div>')
            return
        if parsed.path == "/records":
            rows = records_from_workbook()
            table = "<table>" + "".join(
                "<tr>" + "".join(f"<th>{html.escape(c)}</th>" if i == 0 else f"<td>{html.escape(c)}</td>" for c in row) + "</tr>"
                for i, row in enumerate(rows)
            ) + "</table>"
            self.send_html(f"<h1>外网爆款视频库</h1><p><code>{XLSX_PATH}</code></p>{table}")
            return

        rows = search_urls()
        table_rows = []
        for platform, keyword, url in rows:
            table_rows.append(
                f"<tr><td>{html.escape(platform)}</td><td>{html.escape(keyword)}</td>"
                f'<td><a href="{html.escape(url)}" target="_blank">打开搜索</a></td></tr>'
            )
        body = f"""
<h1>外网爆款视频寻找工具</h1>
<p>先打开搜索链接，看到适合的 TikTok / Instagram / YouTube 视频后，点击 Chrome 插件“采集当前视频”。</p>
<p><a class="button" href="/records" target="_blank">查看已采集表格</a></p>
<p>本地表格：<code>{html.escape(str(XLSX_PATH))}</code></p>
<table><tr><th>Platform</th><th>Keyword / Trend</th><th>Search Link</th></tr>{''.join(table_rows)}</table>
"""
        self.send_html(body)

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path != "/bulk":
            self.send_response(404)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json.dumps({"ok": False, "error": "unknown path"}, ensure_ascii=False).encode("utf-8"))
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            raw = self.rfile.read(length).decode("utf-8", errors="replace")
            payload = json.loads(raw or "{}")
            items = payload.get("items") or []
            saved = 0
            duplicates = 0
            failed = 0
            errors = []
            for item in items:
                if not isinstance(item, dict):
                    failed += 1
                    continue
                params = {str(k): "" if v is None else str(v) for k, v in item.items()}
                try:
                    _, data = save_record(params)
                    if data.get("status") == "duplicate":
                        duplicates += 1
                    else:
                        saved += 1
                except Exception as exc:
                    failed += 1
                    errors.append(str(exc))
            self.send_json(
                {
                    "ok": True,
                    "saved": saved,
                    "duplicates": duplicates,
                    "failed": failed,
                    "errors": errors[:5],
                }
            )
        except Exception as exc:
            self.send_response(500)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False).encode("utf-8"))

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))


def main() -> None:
    if "--export-only" in sys.argv:
        ensure_workbook()
        write_csv_and_md()
        print(MD_PATH)
        return
    ensure_workbook()
    write_csv_and_md()
    url = f"http://127.0.0.1:{PORT}"
    print(f"Viral video finder running: {url}")
    print(f"Workbook: {XLSX_PATH}")
    webbrowser.open(url)
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
