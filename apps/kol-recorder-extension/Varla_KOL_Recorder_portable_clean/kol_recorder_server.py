from __future__ import annotations

import html
import re
import sys
import urllib.parse
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = BASE_DIR / "outputs" / "varla_kol_tracker.xlsx"
SHEET_NAME = "KOL建联台账"
DEFAULT_FIND_NAME = "YOUR_NAME"
PORT = 8765


HEADERS = {
    "Num": "A",
    "Find Name": "B",
    "Date": "C",
    "KOL Name": "D",
    "Platform": "E",
    "Handle": "F",
    "Profile URL": "G",
    "Email": "H",
    "DM Link": "I",
    "Country": "J",
    "Followers": "K",
    "Avg Views": "L",
    "Best Video Views": "M",
    "Content Type": "N",
    "Top Countries": "O",
    "Age Range": "P",
    "Contact Status": "Q",
    "Cooperation Type": "R",
    "Quote": "S",
    "Expected Deliverables": "T",
    "Internal Rating": "U",
    "Next Follow-up Date": "V",
    "Notes": "W",
}


def today_text() -> str:
    return datetime.now().strftime("%Y.%m.%d")


def compact_metric(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(r"[KMBkmb]", text):
        return text.upper()
    cleaned = text.replace(",", "")
    if not re.fullmatch(r"\d+(?:\.\d+)?", cleaned):
        return text
    number = float(cleaned)
    if number >= 1_000_000:
        return f"{number / 1_000_000:.1f}".rstrip("0").rstrip(".") + "M"
    if number >= 1_000:
        decimals = 0 if number >= 100_000 else 1
        return f"{number / 1_000:.{decimals}f}".rstrip("0").rstrip(".") + "K"
    return str(int(number)) if number.is_integer() else str(number)


def normalize_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        return ""
    parsed = urllib.parse.urlparse(url)
    return urllib.parse.urlunparse((parsed.scheme, parsed.netloc, parsed.path.rstrip("/"), "", "", ""))


def infer_platform(url: str) -> str:
    u = url.lower()
    if "tiktok.com" in u:
        return "TK"
    if "instagram.com" in u:
        return "Instagram"
    return ""


def is_valid_profile_url(url: str) -> bool:
    u = normalize_url(url).lower()
    if "127.0.0.1" in u or "localhost" in u:
        return False
    if "tiktok.com/search" in u or "/tag/" in u or "/music/" in u:
        return False
    if re.search(r"tiktok\.com/@[^/?#]+$", u):
        return True
    if re.search(r"instagram\.com/[^/?#]+$", u):
        username = u.rstrip("/").split("/")[-1]
        return username not in {"p", "reel", "reels", "stories", "explore", "accounts", "direct"}
    return False


def infer_handle(url: str, title: str = "") -> str:
    url = normalize_url(url)
    patterns = [
        r"tiktok\.com/@([^/?#]+)",
        r"instagram\.com/([^/?#]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, url, re.I)
        if match:
            handle = match.group(1).strip()
            if handle and handle not in {"p", "reel", "reels", "stories", "explore"}:
                return "@" + handle.lstrip("@")
    match = re.search(r"@([A-Za-z0-9._-]+)", title or "")
    if match:
        return "@" + match.group(1)
    return ""


def infer_name(title: str, handle: str) -> str:
    text = (title or "").strip()
    if not text:
        return handle
    text = re.sub(r"\s*\|\s*TikTok.*$", "", text, flags=re.I)
    text = re.sub(r"\s*•\s*Instagram.*$", "", text, flags=re.I)
    text = re.sub(r"\s*\(@[^)]*\)\s*$", "", text).strip()
    return text[:80] or handle


def find_row(ws, profile_url: str, handle: str) -> int:
    for row in range(3, ws.max_row + 1):
        existing_url = normalize_url(str(ws[f"{HEADERS['Profile URL']}{row}"].value or ""))
        existing_handle = str(ws[f"{HEADERS['Handle']}{row}"].value or "").strip().lower()
        if profile_url and existing_url == profile_url:
            return row
        if handle and existing_handle == handle.lower():
            return row
    for row in range(3, ws.max_row + 1):
        if not ws[f"{HEADERS['Handle']}{row}"].value and not ws[f"{HEADERS['Profile URL']}{row}"].value:
            return row
    return ws.max_row + 1


def save_record(params: dict[str, str]) -> tuple[int, str]:
    url = normalize_url(params.get("url", ""))
    title = params.get("title", "")
    handle = params.get("handle") or infer_handle(url, title)
    platform = params.get("platform") or infer_platform(url)
    kol_name = params.get("kol_name") or infer_name(title, handle)
    find_name = params.get("find_name") or DEFAULT_FIND_NAME

    if not url:
        raise ValueError("缺少当前页面 URL")
    if not is_valid_profile_url(url):
        raise ValueError("当前页面不是 KOL 主页。请回到 TikTok/Instagram 账号主页后再记录，不要在搜索页或记录结果页记录。")
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"找不到 xlsx：{WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)
    ws = wb[SHEET_NAME]
    row = find_row(ws, url, handle)

    ws[f"{HEADERS['Num']}{row}"] = row - 2
    ws[f"{HEADERS['Find Name']}{row}"] = find_name
    ws[f"{HEADERS['Date']}{row}"] = today_text()
    ws[f"{HEADERS['KOL Name']}{row}"] = kol_name
    ws[f"{HEADERS['Platform']}{row}"] = platform
    ws[f"{HEADERS['Handle']}{row}"] = handle
    ws[f"{HEADERS['Profile URL']}{row}"] = url
    ws[f"{HEADERS['DM Link']}{row}"] = url
    ws[f"{HEADERS['Contact Status']}{row}"] = params.get("status") or "待建联"

    optional_fields = {
        "email": "Email",
        "country": "Country",
        "followers": "Followers",
        "avg_views": "Avg Views",
        "best_video_views": "Best Video Views",
        "content_type": "Content Type",
        "top_countries": "Top Countries",
        "age_range": "Age Range",
        "cooperation_type": "Cooperation Type",
        "quote": "Quote",
        "expected_deliverables": "Expected Deliverables",
        "internal_rating": "Internal Rating",
        "next_follow_up": "Next Follow-up Date",
        "notes": "Notes",
    }
    for key, header in optional_fields.items():
        value = params.get(key)
        if value:
            if key in {"followers", "avg_views", "best_video_views"}:
                value = compact_metric(value)
            ws[f"{HEADERS[header]}{row}"] = value

    try:
        wb.save(WORKBOOK_PATH)
    except PermissionError as exc:
        raise PermissionError("xlsx 可能正在被 Excel/飞书占用，请先关闭文件再记录。") from exc
    return row, handle


def page(body: str) -> bytes:
    return f"""<!doctype html>
<meta charset="utf-8">
<title>KOL Recorder</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;max-width:760px;margin:40px auto;line-height:1.6;color:#1f2329}}
input,select,textarea{{width:100%;box-sizing:border-box;padding:8px;margin:4px 0 12px;border:1px solid #c9ccd6;border-radius:6px}}
button{{padding:9px 16px;border:0;border-radius:6px;background:#245bdb;color:white;font-weight:700}}
.ok{{background:#e8f7ee;border:1px solid #b7ebc6;padding:12px;border-radius:8px}}
.err{{background:#fff1f0;border:1px solid #ffa39e;padding:12px;border-radius:8px}}
code{{background:#f5f6f8;padding:2px 4px;border-radius:4px}}
</style>
{body}
""".encode("utf-8")


class Handler(BaseHTTPRequestHandler):
    def send_html(self, body: str):
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(page(body))

    def send_json(self, data: str):
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(data.encode("utf-8"))

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        params = {k: v[0] for k, v in urllib.parse.parse_qs(parsed.query).items()}
        if parsed.path == "/health":
            self.send_json('{"ok":true,"service":"kol-recorder"}')
            return
        if parsed.path == "/add":
            try:
                row, handle = save_record(params)
                body = f'<div class="ok">已记录到第 {row} 行：{html.escape(handle)}</div><p>可以关闭这个页面，继续私信下一个 KOL。</p>'
            except Exception as exc:
                body = f'<div class="err">记录失败：{html.escape(str(exc))}</div><p>常见原因：xlsx 正在打开、URL 为空、文件路径不存在。</p>'
            self.send_html(body)
            return

        bookmarklet = (
            "javascript:(()=>{"
            "const u=encodeURIComponent(location.href);"
            "const t=encodeURIComponent(document.title);"
            f"window.open('http://127.0.0.1:{PORT}/add?url='+u+'&title='+t,'_blank');"
            "})()"
        )
        body = f"""
<h1>KOL 私信记录器</h1>
<p>运行中。发完私信后，点击收藏夹里的记录按钮即可写入：</p>
<p><code>{html.escape(str(WORKBOOK_PATH))}</code></p>
<h2>收藏夹按钮</h2>
<p>新建一个 Chrome 书签，名称填：<b>记录KOL</b></p>
<p>网址填下面这一整段：</p>
<textarea rows="5" readonly>{html.escape(bookmarklet)}</textarea>
<h2>可自动填写</h2>
<p>Date、Platform、Handle、Profile URL、DM Link、Contact Status。</p>
<h2>建议人工补充</h2>
<p>Email、Followers、Avg Views、Best Video Views、Top Countries、Age Range、Quote。</p>
"""
        self.send_html(body)

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))


def main():
    print(f"KOL recorder running: http://127.0.0.1:{PORT}")
    print(f"Workbook: {WORKBOOK_PATH}")
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
