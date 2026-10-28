from __future__ import annotations

import html
import os
import re
import sys
import urllib.parse
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_NAME = "varla_kol_tracker.xlsx"
LOCAL_WORKBOOK_PATH = BASE_DIR / "outputs" / WORKBOOK_NAME
DEFAULT_FIND_NAME = "YOUR_NAME"
PREFERRED_SHEET_NAME = "KOL建联台账"
PORT = 8765

HEADERS = {
    "Num": "A",
    "Date": "B",
    "KOL Name": "C",
    "Platform": "D",
    "Handle": "E",
    "Profile URL": "F",
    "Email": "G",
    "Country": "H",
    "Followers": "I",
    "Avg Views": "J",
    "Best Video Views": "K",
    "Contact Status": "L",
    "Cooperation Type": "M",
    "Quote": "N",
    "Internal Rating": "O",
    "Next Follow-up Date": "P",
}


def today_text() -> str:
    return datetime.now().strftime("%Y.%m.%d")


def compact_metric(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(r"[KMBkmb]", text):
        return text.upper()
    cleaned = text.replace(",", "")
    if not re.fullmatch(r"\d+(?:\.\d+)?", cleaned):
        return text
    number = float(cleaned)
    if number >= 1_000_000:
        return f"{number / 1_000_000:.1f}".rstrip("0").rstrip(".") + "M"
    if number >= 1_000:
        decimals = 0 if number >= 100_000 else 1
        return f"{number / 1_000:.{decimals}f}".rstrip("0").rstrip(".") + "K"
    return str(int(number)) if number.is_integer() else str(number)


def normalize_url(url: str) -> str:
    url = (url or "").strip()
    if not url:
        return ""
    parsed = urllib.parse.urlparse(url)
    return urllib.parse.urlunparse((parsed.scheme, parsed.netloc, parsed.path.rstrip("/"), "", "", ""))


def candidate_workbook_paths() -> list[Path]:
    return [
        LOCAL_WORKBOOK_PATH,
        BASE_DIR / "Varla_KOL_Recorder_portable_clean" / "outputs" / WORKBOOK_NAME,
        BASE_DIR / "Varla_KOL_Recorder_Chrome便携包_文件夹" / "outputs" / WORKBOOK_NAME,
        BASE_DIR.parent / "公共脚本和缓存" / "outputs" / WORKBOOK_NAME,
        BASE_DIR.parent.parent / "outputs" / WORKBOOK_NAME,
    ]


def resolve_workbook_path() -> Path:
    if LOCAL_WORKBOOK_PATH.exists():
        return LOCAL_WORKBOOK_PATH

    for candidate in candidate_workbook_paths()[1:]:
        if candidate.exists():
            LOCAL_WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
            copy2(candidate, LOCAL_WORKBOOK_PATH)
            return LOCAL_WORKBOOK_PATH

    raise FileNotFoundError("找不到 varla_kol_tracker.xlsx，请确认本地台账文件存在。")


def resolve_sheet_name(workbook) -> str:
    return PREFERRED_SHEET_NAME if PREFERRED_SHEET_NAME in workbook.sheetnames else workbook.sheetnames[0]


def infer_platform(url: str) -> str:
    lowered = (url or "").lower()
    if "tiktok.com" in lowered:
        return "TK"
    if "instagram.com" in lowered:
        return "Instagram"
    return ""


def is_valid_profile_url(url: str) -> bool:
    lowered = normalize_url(url).lower()
    if not lowered or "127.0.0.1" in lowered or "localhost" in lowered:
        return False
    if any(fragment in lowered for fragment in ("/search", "/explore", "/tag/", "/music/", "/reel/", "/video/")):
        return False
    if re.search(r"tiktok\.com/@[^/?#]+$", lowered):
        return True
    if re.search(r"instagram\.com/[^/?#]+$", lowered):
        username = lowered.rstrip("/").split("/")[-1]
        return username not in {"p", "reel", "reels", "stories", "explore", "accounts", "direct"}
    return False


def infer_handle(url: str, title: str = "") -> str:
    normalized = normalize_url(url)
    for pattern in (r"tiktok\.com/@([^/?#]+)", r"instagram\.com/([^/?#]+)"):
        match = re.search(pattern, normalized, re.I)
        if match:
            return "@" + match.group(1).lstrip("@")
    title_match = re.search(r"@([A-Za-z0-9._-]+)", title or "")
    return f"@{title_match.group(1)}" if title_match else ""


def infer_name(title: str, handle: str) -> str:
    text = (title or "").strip()
    if not text or text.lower() == "instagram":
        return handle
    text = re.sub(r"\s*\|\s*TikTok.*$", "", text, flags=re.I)
    text = re.sub(r"\s*[-·•|]\s*Instagram.*$", "", text, flags=re.I)
    text = re.sub(r"\s*\(@[^)]*\)\s*$", "", text).strip()
    return (text[:80] or handle).strip()


def find_row(worksheet, profile_url: str, handle: str) -> int:
    for row in range(3, worksheet.max_row + 1):
        existing_url = normalize_url(str(worksheet[f"{HEADERS['Profile URL']}{row}"].value or ""))
        existing_handle = str(worksheet[f"{HEADERS['Handle']}{row}"].value or "").strip().lower()
        if profile_url and existing_url == profile_url:
            return row
        if handle and existing_handle == handle.lower():
            return row
    for row in range(3, worksheet.max_row + 1):
        if not worksheet[f"{HEADERS['Handle']}{row}"].value and not worksheet[f"{HEADERS['Profile URL']}{row}"].value:
            return row
    return max(3, worksheet.max_row + 1)


def latest_records(limit: int = 8) -> list[dict[str, str]]:
    workbook_path = resolve_workbook_path()
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[resolve_sheet_name(workbook)]
    records: list[dict[str, str]] = []
    for row in range(worksheet.max_row, 2, -1):
        handle = str(worksheet[f"{HEADERS['Handle']}{row}"].value or "").strip()
        profile_url = str(worksheet[f"{HEADERS['Profile URL']}{row}"].value or "").strip()
        if not handle and not profile_url:
            continue
        records.append(
            {
                "num": str(worksheet[f"{HEADERS['Num']}{row}"].value or ""),
                "date": str(worksheet[f"{HEADERS['Date']}{row}"].value or ""),
                "name": str(worksheet[f"{HEADERS['KOL Name']}{row}"].value or ""),
                "platform": str(worksheet[f"{HEADERS['Platform']}{row}"].value or ""),
                "handle": handle,
                "country": str(worksheet[f"{HEADERS['Country']}{row}"].value or ""),
                "followers": str(worksheet[f"{HEADERS['Followers']}{row}"].value or ""),
                "avg_views": str(worksheet[f"{HEADERS['Avg Views']}{row}"].value or ""),
                "status": str(worksheet[f"{HEADERS['Contact Status']}{row}"].value or ""),
                "profile_url": profile_url,
            }
        )
        if len(records) >= limit:
            break
    return records


def recent_table_html(limit: int = 8) -> str:
    rows = latest_records(limit)
    if not rows:
        return '<p class="muted">还没有已记录内容。</p>'

    cells = []
    for item in rows:
        link = html.escape(item["profile_url"])
        cells.append(
            "<tr>"
            f"<td>{html.escape(item['num'])}</td>"
            f"<td>{html.escape(item['date'])}</td>"
            f"<td>{html.escape(item['platform'])}</td>"
            f"<td>{html.escape(item['handle'])}</td>"
            f"<td>{html.escape(item['name'])}</td>"
            f"<td>{html.escape(item['country'])}</td>"
            f"<td>{html.escape(item['followers'])}</td>"
            f"<td>{html.escape(item['avg_views'])}</td>"
            f"<td>{html.escape(item['status'])}</td>"
            f'<td><a href="{link}" target="_blank">打开主页</a></td>'
            "</tr>"
        )

    return (
        '<div class="table-wrap"><table>'
        "<thead><tr>"
        "<th>序号</th><th>Date</th><th>Platform</th><th>Handle</th><th>KOL Name</th>"
        "<th>Country</th><th>Followers</th><th>Avg Views</th><th>Status</th><th>Profile</th>"
        "</tr></thead>"
        f"<tbody>{''.join(cells)}</tbody>"
        "</table></div>"
    )


def open_workbook() -> Path:
    workbook_path = resolve_workbook_path()
    os.startfile(str(workbook_path))
    return workbook_path


def save_record(params: dict[str, str]) -> tuple[int, str, Path]:
    url = normalize_url(params.get("url", ""))
    title = params.get("title", "")
    handle = params.get("handle") or infer_handle(url, title)
    platform = params.get("platform") or infer_platform(url)
    kol_name = params.get("kol_name") or infer_name(title, handle)
    if not kol_name or kol_name.lower() == "instagram":
        kol_name = handle
    if not url:
        raise ValueError("缺少当前主页 URL")
    if not is_valid_profile_url(url):
        raise ValueError("当前页面不是 KOL 主页，请回到 TikTok / Instagram 账号主页后再记录。")

    workbook_path = resolve_workbook_path()
    workbook = load_workbook(workbook_path)
    worksheet = workbook[resolve_sheet_name(workbook)]
    row = find_row(worksheet, url, handle)

    worksheet[f"{HEADERS['Num']}{row}"] = row - 2
    worksheet[f"{HEADERS['Date']}{row}"] = today_text()
    worksheet[f"{HEADERS['KOL Name']}{row}"] = kol_name
    worksheet[f"{HEADERS['Platform']}{row}"] = platform
    worksheet[f"{HEADERS['Handle']}{row}"] = handle
    worksheet[f"{HEADERS['Profile URL']}{row}"] = url
    worksheet[f"{HEADERS['Contact Status']}{row}"] = params.get("status") or "待建联"

    optional_fields = {
        "email": "Email",
        "country": "Country",
        "followers": "Followers",
        "avg_views": "Avg Views",
        "best_video_views": "Best Video Views",
        "top_countries": "Top Countries",
        "age_range": "Age Range",
        "cooperation_type": "Cooperation Type",
        "quote": "Quote",
        "internal_rating": "Internal Rating",
        "next_follow_up": "Next Follow-up Date",
    }
    for key, header in optional_fields.items():
        value = params.get(key)
        if value:
            if header not in HEADERS:
                continue
            if key in {"followers", "avg_views", "best_video_views"}:
                value = compact_metric(value)
            worksheet[f"{HEADERS[header]}{row}"] = value

    try:
        workbook.save(workbook_path)
    except PermissionError as exc:
        raise PermissionError("xlsx 可能正在被 Excel 或飞书占用，请先关闭文件后再记录。") from exc

    return row, handle, workbook_path


def page(body: str) -> bytes:
    return f"""<!doctype html>
<meta charset="utf-8">
<title>KOL Recorder</title>
<style>
body{{font-family:Arial,'Microsoft YaHei',sans-serif;max-width:1080px;margin:40px auto;line-height:1.6;color:#1f2329;padding:0 20px}}
.panel{{background:#f7f8fa;border:1px solid #e5e6eb;border-radius:14px;padding:18px 20px;margin:18px 0}}
.button-row{{display:flex;gap:12px;flex-wrap:wrap;margin:12px 0 0}}
.button-link{{padding:11px 18px;border:0;border-radius:10px;background:#245bdb;color:white;font-weight:700;font-size:15px;text-decoration:none;display:inline-block}}
.button-link.secondary{{background:#344054}}
.ok{{background:#e8f7ee;border:1px solid #b7ebc6;padding:12px;border-radius:8px}}
.err{{background:#fff1f0;border:1px solid #ffa39e;padding:12px;border-radius:8px}}
.muted{{color:#646a73}}
code{{background:#f5f6f8;padding:2px 4px;border-radius:4px;word-break:break-all}}
textarea{{width:100%;box-sizing:border-box;padding:10px 12px;border:1px solid #c9ccd6;border-radius:8px;font-size:14px}}
summary{{cursor:pointer;font-weight:700}}
.table-wrap{{overflow:auto;border:1px solid #e5e6eb;border-radius:12px;background:white}}
table{{width:100%;border-collapse:collapse;font-size:14px}}
th,td{{padding:10px 12px;border-bottom:1px solid #edf0f2;text-align:left;vertical-align:top;white-space:nowrap}}
th{{background:#f7f8fa;font-weight:700}}
tr:last-child td{{border-bottom:0}}
a{{color:#245bdb;text-decoration:none}}
a:hover{{text-decoration:underline}}
</style>
{body}
""".encode("utf-8")


def render_home() -> str:
    workbook_path = resolve_workbook_path()
    bookmarklet = (
        "javascript:(()=>{"
        "const u=encodeURIComponent(location.href);"
        "const t=encodeURIComponent(document.title);"
        f"window.open('http://127.0.0.1:{PORT}/add?url='+u+'&title='+t,'_blank');"
        "})()"
    )
    return f"""
<h1>KOL 私信记录器</h1>
<p class="muted">正常使用时，你只需要在 TikTok / Instagram 页面右侧点一次 <strong>记录当前KOL</strong>，就会直接同步到本地 xlsx。</p>

<div class="panel">
  <p><strong>当前台账：</strong><code>{html.escape(str(workbook_path))}</code></p>
  <div class="button-row">
    <a class="button-link" href="/open-workbook">打开本地 XLSX</a>
    <a class="button-link secondary" href="/" target="_self">刷新状态</a>
  </div>
</div>

<div class="panel">
  <h2>怎么用</h2>
  <p>1. 打开 TikTok / Instagram 的 KOL 主页。</p>
  <p>2. 点页面右侧悬浮窗里的 <strong>记录当前KOL</strong>。</p>
  <p>3. 数据会直接写入上面这份本地 xlsx。</p>
  <details>
    <summary>旧版书签脚本，仅兜底</summary>
    <p class="muted">如果悬浮按钮临时失效，再用这个收藏夹脚本。平时不用管它。</p>
    <textarea rows="4" readonly>{html.escape(bookmarklet)}</textarea>
  </details>
</div>

<div class="panel">
  <h2>最近已记录</h2>
  {recent_table_html(8)}
</div>
"""


class Handler(BaseHTTPRequestHandler):
    def send_html(self, body: str):
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(page(body))

    def send_json(self, data: str):
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(data.encode("utf-8"))

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        params = {key: values[0] for key, values in urllib.parse.parse_qs(parsed.query).items()}

        if parsed.path == "/health":
            self.send_json('{"ok":true,"service":"kol-recorder"}')
            return

        if parsed.path == "/open-workbook":
            try:
                workbook_path = open_workbook()
                body = f"""
<div class="ok">已请求打开本地 XLSX：{html.escape(str(workbook_path))}</div>
<div class="button-row">
  <a class="button-link secondary" href="/" target="_self">返回状态页</a>
</div>
"""
            except Exception as exc:
                body = f"""
<div class="err">打开失败：{html.escape(str(exc))}</div>
<div class="button-row">
  <a class="button-link secondary" href="/" target="_self">返回状态页</a>
</div>
<p class="muted">如果 Excel 没弹出，请直接从上面的本地路径打开 xlsx。</p>
"""
            self.send_html(body)
            return

        if parsed.path == "/add":
            try:
                row, handle, workbook_path = save_record(params)
                display_name = handle or params.get("url", "")
                body = f"""
<div class="ok">已记录到第 {row} 行：{html.escape(display_name)}</div>
<div class="button-row">
  <a class="button-link" href="/open-workbook">打开 XLSX</a>
  <a class="button-link secondary" href="/" target="_self">返回状态页</a>
</div>
<p class="muted">已经同步到本地台账，可以继续记录下一个 KOL。</p>
"""
            except Exception as exc:
                body = f"""
<div class="err">记录失败：{html.escape(str(exc))}</div>
<div class="button-row">
  <a class="button-link secondary" href="/" target="_self">返回状态页</a>
</div>
<p class="muted">常见原因：当前不是账号主页，或者 xlsx 正在被 Excel / 飞书占用。</p>
"""
            self.send_html(body)
            return

        self.send_html(render_home())

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))


def main():
    workbook_path = resolve_workbook_path()
    print(f"KOL recorder running: http://127.0.0.1:{PORT}")
    print(f"Workbook: {workbook_path}")
    ThreadingHTTPServer(("127.0.0.1", PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
